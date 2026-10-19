"""Exports DXF, PDF et XLSX : fichiers valides, contenu fidèle aux moteurs, mentions de validation."""
import io
import json
import subprocess

import ezdxf
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.moteurs import chaine
from app.moteurs.programme import Programme
from app.rendu import dxf, pdf
from tests.test_api import _compte  # noqa: F401  (même gabarit de comptes)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path}/test.db")
    monkeypatch.setenv("ADMIN_EMAIL", "admin@exemple.cd")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


@pytest.fixture
def comptes(client):
    admin = _compte(client, "admin@exemple.cd", "Admin")
    return {"admin": admin, "c": _compte(client, "c@exemple.cd", "Constructeur"),
            "ing": _compte(client, "ing@exemple.cd", "Ir. Mukendi", role="ingenieur", admin=admin)}


@pytest.fixture
def resultats(programme_ref) -> chaine.Resultats:
    return chaine.calculer(programme_ref.model_copy(update={"terrain_largeur": 18}))


def _doc_dxf(resultats) -> ezdxf.document.Drawing:
    octets = dxf.exporter(resultats.implantation, nom_projet="Villa", elec=resultats.electricite.model_dump(),
                          structure=resultats.structure.model_dump())
    return ezdxf.read(io.StringIO(octets.decode("utf-8")))


def test_dxf_valide_sans_erreur(resultats):
    doc = _doc_dxf(resultats)
    audit = doc.audit()
    assert not audit.errors and not audit.fixes
    for calque in ("MURS", "PORTES", "FENETRES", "COTES", "TEXTES", "ELEC", "STRUCTURE"):
        assert doc.layers.has_entry(calque), calque
    assert doc.units == ezdxf.units.CM


def test_dxf_contenu_fidele(resultats):
    doc = _doc_dxf(resultats)
    msp = doc.modelspace()
    par_calque = {c: [e for e in msp if e.dxf.layer == c] for c in ("MURS", "TEXTES", "ELEC", "STRUCTURE", "PORTES")}
    impl = resultats.implantation
    assert len(par_calque["MURS"]) == len(impl.murs)
    # Un arc de débattement par porte battante
    battantes = [o for o in impl.ouvertures if o.type == "porte" and o.style != "basculante"]
    assert len([e for e in par_calque["PORTES"] if e.dxftype() == "ARC"]) == len(battantes)
    textes = " ".join(e.dxf.text for e in par_calque["TEXTES"] if e.dxftype() == "TEXT")
    for pc in impl.pieces:
        if not pc.exterieure:
            assert pc.nom in textes
    assert "À valider par un ingénieur qualifié" in textes
    assert len([e for e in par_calque["ELEC"] if e.dxftype() == "CIRCLE"]) > 20
    assert len([e for e in par_calque["STRUCTURE"] if e.dxftype() == "LWPOLYLINE"]) == (
        len(resultats.structure.semelles) + len(resultats.structure.poteaux))
    # Les murs sont dans l'emprise (axe Y vers le haut : y négatif)
    ys = [p[1] / 100 for e in par_calque["MURS"] for p in e.get_points("xy")]  # cm → m
    xs = [p[0] / 100 for e in par_calque["MURS"] for p in e.get_points("xy")]
    assert max(ys) <= -impl.emprise.y_min + 1e-6 and min(ys) >= -impl.emprise.y_max - 1e-6
    assert min(xs) == pytest.approx(impl.emprise.x_min) and max(xs) == pytest.approx(impl.emprise.x_max)
    # Hauteur de texte et rayon d'arc à l'échelle (cm)
    assert {round(e.dxf.height, 3) for e in par_calque["TEXTES"] if e.dxftype() == "TEXT"} >= {16.0, 13.0}
    assert {round(e.dxf.radius) for e in par_calque["PORTES"] if e.dxftype() == "ARC"} >= {80}


def _texte_pdf(octets: bytes) -> str:
    return subprocess.run(["pdftotext", "-layout", "-", "-"], input=octets, capture_output=True, check=True).stdout.decode()


def test_pdf_cartouche_et_contenu(resultats):
    meta = {"nom": "Villa de référence", "version": 3, "empreinte": "ab" * 32, "modifie_le": "2026-10-18T09:00:00+00:00",
            "version_moteurs": resultats.version_moteurs}
    octets = pdf.rendre(meta, resultats, [])
    assert octets.startswith(b"%PDF-")
    texte = _texte_pdf(octets)
    assert "Villa de référence" in texte and "1:100" in texte and "2026-10-18" in texte and "v3" in texte
    assert "Rez-de-chaussée" in texte and "abababababab" in texte
    assert texte.count("À valider par un ingénieur qualifié") >= 3  # cartouche + pied de page des pages
    assert "document non contractuel" in texte
    for pc in resultats.implantation.pieces:
        if pc.type in ("chambre", "sejour", "cuisine"):
            assert pc.nom in texte
    for c in resultats.electricite.circuits:
        assert f"C{c.num}" in texte
    assert f"{resultats.metre.total:,.2f}".replace(",", " ").replace(".", ",")[-6:] in texte.replace(" ", " ")
    # Aucun symbole perdu (Ω, ≤, σ…) : pas de caractère de remplacement
    assert "�" not in texte and "■" not in texte


def test_pdf_signe_tous_modules(resultats):
    meta = {"nom": "Villa", "version": 1, "empreinte": "cd" * 32, "modifie_le": "2026-10-18T09:00:00+00:00",
            "version_moteurs": resultats.version_moteurs}
    sig = [{"module": m, "nom": "Ir. Mukendi", "numero_ordre": "ONI-0042", "date": "2026-10-18T10:00:00+00:00",
            "commentaire": ""} for m in pdf.MODULES]
    texte = _texte_pdf(pdf.rendre(meta, resultats, sig))
    assert "Signé par Ir. Mukendi (n° ONI-0042)" in texte
    assert "À valider par un ingénieur qualifié" not in texte
    assert "Tous les modules sont signés" in texte


def test_pdf_texte_sans_symbole_perdu():
    assert pdf.ascii_pdf("σ sol ≤ 0,15 MPa · 45 Ω → ΔU") == "sigma sol <= 0,15 MPa · 45 ohm -> DU"
    assert pdf.ascii_pdf("é à ç — ² × °") == "é à ç — ² × °"


def test_pdf_tous_les_textes_des_moteurs_sont_imprimables():
    """Aucun texte produit par les moteurs ne perd de caractère (cp1252) après remplacement."""
    for upd in ({}, {"toiture": "terrasse", "couverture": "dalle beton"}, {"toiture": "4 pans", "couverture": "tuiles"},
                {"terrain_largeur": 18, "garage": True, "bureau": True, "chambres": 4}):
        r = chaine.calculer(Programme(**{"terrain_largeur": 25, **upd}))
        textes = [c.titre + c.detail + c.reference_norme for c in r.controles] + [h.libelle + h.valeur + h.source
                                                                                   for h in r.hypotheses]
        for t in textes:
            sortie = pdf.ascii_pdf(t)
            sortie.encode("cp1252")  # imprimable tel quel
            assert sortie.count("?") == t.count("?"), t  # aucun caractère remplacé par un point d'interrogation


def test_api_exports(client, comptes):
    c = comptes["c"]
    projet = client.post("/api/projets/generer", json={"programme": {"terrain_largeur": 18}}, headers=c).json()
    pid = projet["id"]

    r = client.get(f"/api/projets/{pid}/plan.dxf", headers=c)
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert ezdxf.read(io.StringIO(r.content.decode())).audit().errors == []

    r = client.get(f"/api/projets/{pid}/dossier.pdf", headers=c)
    assert r.status_code == 200 and r.content.startswith(b"%PDF-") and r.headers["content-type"] == "application/pdf"
    assert "À valider par un ingénieur qualifié" in _texte_pdf(r.content)

    r = client.get(f"/api/projets/{pid}/metre.xlsx", headers=c)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Métré", "Quantités", "Hypothèses"]
    assert "document non contractuel" in wb["Métré"]["A3"].value

    for ext in ("plan.dxf", "dossier.pdf", "metre.xlsx"):
        assert client.get(f"/api/projets/{pid}/{ext}").status_code == 401
        assert client.get(f"/api/projets/{pid}/{ext}", headers=comptes["admin"]).status_code == 200
    autre = _compte(client, "autre@exemple.cd", "Autre")
    assert client.get(f"/api/projets/{pid}/dossier.pdf", headers=autre).status_code == 403


def test_xlsx_formules_et_totaux_recalcules(client, comptes, tmp_path):
    c = comptes["c"]
    pid = client.post("/api/projets/generer", json={"programme": {"terrain_largeur": 18}}, headers=c).json()["id"]
    contenu = client.get(f"/api/projets/{pid}/metre.xlsx", headers=c).content
    projet = client.get(f"/api/projets/{pid}", headers=c).json()
    metre = projet["resultats"]["metre"]
    wb = load_workbook(io.BytesIO(contenu))
    ws = wb["Métré"]
    formules = [cell.value for row in ws.iter_rows(min_row=6) for cell in row if isinstance(cell.value, str)
                and cell.value.startswith("=")]
    assert any(f.startswith("=D") for f in formules) and any(f.startswith("=SUM(F") for f in formules)
    # Recalcul réel par LibreOffice : le total du classeur égale celui du moteur
    fichier = tmp_path / "metre.xlsx"
    fichier.write_bytes(contenu)
    try:
        subprocess.run(["soffice", "--headless", "--convert-to", "csv", "--outdir", str(tmp_path), str(fichier)],
                       check=True, capture_output=True, timeout=120)
    except (OSError, subprocess.SubprocessError):
        pytest.skip("LibreOffice indisponible pour recalculer le classeur")
    lignes = (tmp_path / "metre.csv").read_text(encoding="utf-8").splitlines()
    total = next(ligne for ligne in lignes if "Total estimatif" in ligne)
    valeur = float(total.rstrip(",").split(",")[-1].replace('"', "").replace(" ", "").replace(" ", ""))
    assert valeur == pytest.approx(round(metre["total"], 2), abs=0.02)
    assert json.dumps(metre["devise"])  # devise portée par l'en-tête de colonne
