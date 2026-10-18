"""Export XLSX du métré (openpyxl) : totaux par formules, mention « à mettre à jour » et « à valider »."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.moteurs import chaine

MENTION = "À valider par un ingénieur qualifié"
BLEU = "1F5FAE"
FORMAT_QTE = '#,##0.00'
FORMAT_ARGENT = '#,##0.00'


def _entete(ws, ligne: int, titres: list[str]) -> None:
    for i, t in enumerate(titres, 1):
        c = ws.cell(row=ligne, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLEU)
        c.alignment = Alignment(horizontal="left" if i <= 2 else "right" if i != 3 else "center", vertical="center",
                                wrap_text=True)


def rendre(nom_projet: str, res: chaine.Resultats, signe_tout: bool) -> bytes:
    me = res.metre
    wb = Workbook()
    ws = wb.active
    ws.title = "Métré"
    ws["A1"] = f"Métré et devis estimatif — {nom_projet}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Prix {me.mention_prix}"
    ws["A2"].font = Font(italic=True, color="A85B00")
    ws["A3"] = "Tous les modules sont signés par un ingénieur." if signe_tout else f"{MENTION} — document non contractuel."
    ws["A3"].font = Font(bold=True, color="1B7A3D" if signe_tout else "B3261E")
    _entete(ws, 5, ["Lot", "Poste", "Unité", "Quantité", f"Prix unitaire ({me.devise})", f"Total ({me.devise})"])
    ligne = 6
    premiere_ligne = ligne
    sous_totaux: list[int] = []
    bord = Border(bottom=Side(style="hair", color="999999"))
    for lot in dict.fromkeys(x.lot for x in me.lignes):
        debut = ligne
        for x in (y for y in me.lignes if y.lot == lot):
            ws.cell(row=ligne, column=1, value=lot)
            ws.cell(row=ligne, column=2, value=x.poste)
            ws.cell(row=ligne, column=3, value=x.unite).alignment = Alignment(horizontal="center")
            ws.cell(row=ligne, column=4, value=round(x.quantite, 4)).number_format = FORMAT_QTE
            ws.cell(row=ligne, column=5, value=x.prix_unitaire).number_format = FORMAT_ARGENT
            ws.cell(row=ligne, column=5).font = Font(color="0000FF")  # donnée saisie : modifiable
            ws.cell(row=ligne, column=6, value=f"=D{ligne}*E{ligne}").number_format = FORMAT_ARGENT
            for col in range(1, 7):
                ws.cell(row=ligne, column=col).border = bord
            ligne += 1
        ws.cell(row=ligne, column=2, value=f"Sous-total {lot}").font = Font(bold=True)
        c = ws.cell(row=ligne, column=6, value=f"=SUM(F{debut}:F{ligne - 1})")
        c.number_format, c.font = FORMAT_ARGENT, Font(bold=True)
        for col in range(1, 7):
            ws.cell(row=ligne, column=col).fill = PatternFill("solid", fgColor="EEF2F7")
        sous_totaux.append(ligne)
        ligne += 1
    ws.cell(row=ligne + 1, column=2, value="Total estimatif").font = Font(bold=True, size=12)
    total = ws.cell(row=ligne + 1, column=6, value="=" + "+".join(f"F{r}" for r in sous_totaux))
    total.number_format, total.font = FORMAT_ARGENT, Font(bold=True, size=12)
    total.border = Border(top=Side(style="medium"), bottom=Side(style="double"))
    for col, largeur in zip("ABCDEF", (16, 58, 9, 13, 20, 18), strict=True):
        ws.column_dimensions[col].width = largeur
    ws.freeze_panes = ws.cell(row=premiere_ligne, column=1)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "5:5"

    q = wb.create_sheet("Quantités")
    _entete(q, 1, ["Grandeur", "", "", "Valeur", "", ""])
    for i, (k, v) in enumerate(sorted(me.quantites.items()), 2):
        q.cell(row=i, column=1, value=k)
        q.cell(row=i, column=4, value=round(v, 4)).number_format = FORMAT_QTE
    q.column_dimensions["A"].width = 28
    q.column_dimensions["D"].width = 14

    h = wb.create_sheet("Hypothèses")
    for i, t in enumerate(["Module", "Hypothèse", "Valeur", "Source"], 1):
        c = h.cell(row=1, column=i, value=t)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=BLEU)
    for i, x in enumerate(res.hypotheses, 2):
        for j, v in enumerate((x.module, x.libelle, x.valeur, x.source), 1):
            h.cell(row=i, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    for col, largeur in zip("ABCD", (16, 44, 36, 56), strict=True):
        h.column_dimensions[get_column_letter("ABCD".index(col) + 1)].width = largeur

    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = f"Métré — {nom_projet}"
    wb.properties.creator = "Assistant Plan Ingénierie"
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()
